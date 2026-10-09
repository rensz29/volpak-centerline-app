"""Sample OCAPs for the tests (ADR-0031), made here: the plant's real ones aren't in the repository. And a stand-in
for clamd.

- `vertical_pdf()`: OCAP-017, English, four pages, with a header and a "Page n of m" footer on each, numbered headings
  set bold and larger, and the heater-fault section on page 3.
- `nozzle_pdf()`: OCAP-021, Filipino, two pages.
- `bottom_docx()`: OCAP-030, a Word file with heading styles, a table, and a page break before its last section.
- `sealer_xlsx()`: OCAP-040, an Excel workbook laid out like the plant's (ADR-0039): an OCAP sheet with a title, a
  group-header row and rows merged into blocks; a troubleshooting sheet with a cause column and a note below its table;
  and a hidden sheet.
- `infected_docx()`: that Word file with the EICAR test file hidden inside it, which clamd finds (EICAR only counts at
  the start of a file, so one written after a PDF header passes a real clamd).
- `FakeClamd`: clamd's INSTREAM protocol on a free local port.
"""

from __future__ import annotations

import io
import socketserver
import struct
import threading
import zipfile
from datetime import datetime, timezone

from centerline_api.settings import ScannerSettings

EICAR = rb"X5O!P%@AP[4\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"  # the harmless antivirus test file

VERTICAL = [
    (1, "OCAP-017: Vertical sealing temperature off its centerline", 16, None),
    (1, "1 Purpose", 13, "What the line does when a vertical heater's HMI setpoint or actual temperature leaves "
                         "its centerline. It applies to Vertical 1 to Vertical 6 on the Volpak filler."),
    (1, "2 Safety", 13, "The jaws are hot. Lock out the machine before touching a heater or a thermocouple. "
                        "Wear heat-resistant gloves."),
    (2, "3 Corrective actions", 13, None),
    (2, "3.1 HMI setpoint changed by mistake", 11,
     "If the HMI setpoint of a vertical heater differs from its centerline target, set it back to the target on "
     "the HMI. Check with the shift leader whether a new film roll or a trial asked for the change. Record the "
     "reason in Centerline."),
    (3, "3.2 Heater element or thermocouple fault", 11,
     "If the actual temperature stays below the setpoint while the HMI setpoint is on target, the heater element "
     "or its thermocouple may have failed. Check the heater's current, then the thermocouple's reading against a "
     "handheld probe. Replace the faulty part and test-seal ten sachets before releasing the line."),
    (4, "4 Escalation", 13, "Call the maintenance technician when the temperature doesn't come back within "
                            "15 minutes. Tell the Manager on shift about any product sealed outside the centerline."),
]

NOZZLE = [
    (1, "OCAP-021: Presyon ng nozzle na wala sa centerline", 16, None),
    (1, "1 Layunin", 13, "Ang gagawin kapag ang presyon ng nozzle ay mas mataas o mas mababa sa centerline. "
                         "Para ito sa lahat ng nozzle ng Volpak filler."),
    (2, "2 Mga hakbang kapag barado ang nozzle", 13,
     "Kapag tumaas ang presyon at bumaba ang dami ng laman ng sachet, maaaring barado ang nozzle. Itigil ang "
     "makina, linisin ang nozzle, at suriin ang filter bago muling patakbuhin."),
]


def _pdf(code: str, title: str, items: list) -> bytes:
    from fpdf import FPDF

    pages = max(p for p, *_ in items)

    class Ocap(FPDF):
        def header(self):
            self.set_font("helvetica", size=8)
            self.cell(0, 6, f"Volpak filler - {code} - {title}", new_x="LMARGIN", new_y="NEXT")
            self.ln(4)

        def footer(self):
            self.set_y(-15)
            self.set_font("helvetica", size=8)
            self.cell(0, 6, f"Page {self.page_no()} of {pages}", align="C")

    pdf = Ocap()
    pdf.set_creation_date(datetime(2026, 1, 1, tzinfo=timezone.utc))  # the same bytes every time
    pdf.set_auto_page_break(True, margin=20)
    for page, heading, size, body in items:
        while pdf.page_no() < page:
            pdf.add_page()
        pdf.set_font("helvetica", style="B", size=size)
        pdf.multi_cell(0, size * 0.5, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
        if body:
            pdf.set_font("helvetica", size=10)
            pdf.multi_cell(0, 5, body, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)
    return bytes(pdf.output())


def vertical_pdf() -> bytes:
    return _pdf("OCAP-017", "Vertical sealing temperature", VERTICAL)


def nozzle_pdf() -> bytes:
    return _pdf("OCAP-021", "Presyon ng nozzle", NOZZLE)


def bottom_docx() -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK

    doc = Document()
    doc.add_heading("OCAP-030: Bottom sealing temperature", level=0)
    doc.add_heading("Purpose", level=1)
    doc.add_paragraph("What to do when the front or rear bottom sealing temperature leaves its centerline.")
    doc.add_heading("Checks", level=1)
    doc.add_paragraph("Compare the HMI setpoint with the centerline target for the bottom heater.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Zone", "Target"
    table.cell(1, 0).text, table.cell(1, 1).text = "Rear bottom", "180 °C"
    run = doc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)
    doc.add_heading("Bottom heater fault", level=2)
    doc.add_paragraph("If the rear bottom actual temperature drops while its setpoint is on target, check the "
                      "heater cartridge and its fuse.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def infected_docx() -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(bottom_docx())) as z, zipfile.ZipFile(out, "w") as w:
        for name in z.namelist():
            w.writestr(name, z.read(name))
        w.writestr("word/media/eicar.com", EICAR)
    return out.getvalue()


class FakeClamd:
    """clamd's INSTREAM: it answers FOUND for the EICAR test string, OK otherwise, or what `reply` says."""

    def __init__(self):
        self.scanned: list[bytes] = []
        self.reply: bytes | None = None
        fake = self

        class Handler(socketserver.BaseRequestHandler):
            def handle(self):
                f = self.request.makefile("rb")
                command = b""
                while not command.endswith(b"\0"):
                    command += f.read(1)
                data = b""
                while (n := struct.unpack(">I", f.read(4))[0]) != 0:
                    data += f.read(n)
                fake.scanned.append(data)
                self.request.sendall(fake.reply or (b"stream: Eicar-Signature FOUND\0" if b"EICAR-STANDARD-ANTIVIRUS" in data
                                                    else b"stream: OK\0"))

        self.server = socketserver.ThreadingTCPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.settings = ScannerSettings("clamd", "127.0.0.1", self.server.server_address[1], 5)

    def close(self):
        self.server.shutdown()
        self.server.server_close()


# -- Excel (ADR-0039) -------------------------------------------------------------------------------------------------

SEALING = {  # (row, column): text; column 1 is A
    (1, 1): "Out of Control Plan", (2, 2): "Why change?", (2, 5): "If no good",
    (3, 1): "OCAP #", (3, 2): "Phenomena", (3, 3): "Assembly", (3, 4): "Centerline Name", (3, 5): "Inspection",
    (3, 6): "Standard Value",
    (4, 1): 1, (4, 2): "Weak seal on the vertical side", (4, 3): "Vertical Sealer", (4, 4): "Vertical Sealing 1 - 6",
    (4, 5): "1. Compare the vertical heaters on the HMI", (5, 5): "2. Check the seal width", (4, 6): "200 - 220",
    (6, 1): 2, (6, 2): "Weak bottom seal", (6, 3): "Bottom Sealer", (6, 4): "Bottom Sealing", (6, 5): "1. Check the bottom jaw",
    (6, 6): "165 - 190",
    (7, 1): 3, (7, 2): "Weak top seal", (7, 3): "Top Sealer", (7, 4): "Top Sealing", (7, 5): "1. Check the top jaw",
    (7, 6): "185 - 200",
}
TROUBLESHOOTING = {
    (1, 1): "SEALER TROUBLESHOOTING", (2, 1): "Applies to the top, bottom and vertical sealers",
    (3, 1): "No.", (3, 2): "Sealer / Area", (3, 3): "Possible Phenomenon", (3, 4): "Possible Cause", (3, 5): "Operator Mitigation",
    (4, 1): 1, (4, 2): "Top / Bottom / Vertical", (4, 3): "Seal separates easily", (4, 4): "Temperature too low; heater not warm yet",
    (4, 5): "Hold the pouches and let the heaters reach their preset.",
    (5, 1): 2, (5, 2): "Top / Bottom / Vertical", (5, 3): "Burnt or brittle seal", (5, 4): "Temperature too high",
    (5, 5): "Return to the centerline and let it settle.",
    (6, 1): 3, (6, 2): "Cutter", (6, 3): "Seal not centred on the cutter", (6, 4): "Film position",
    (6, 5): "Centre the sealing plate.",
    (7, 1): 4, (7, 2): "Vertical Sealer", (7, 3): "Vertical seal narrower than 6 mm", (7, 4): "Misalignment",
    (7, 5): "Check the jaw contact.",
    (9, 1): "CONTROL NOTE: tell the Cell Lead before running outside the centerline.",
}


def _col(c: int) -> str:
    name = ""
    while c:
        c, rest = divmod(c - 1, 26)
        name = chr(65 + rest) + name
    return name


def xlsx(sheets: list[tuple[str, dict, list[str], bool]]) -> bytes:
    """A workbook from (name, cells, merged ranges, hidden) sheets: text goes to the shared strings, numbers stay numbers."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    strings: list[str] = []

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i, (_, cells, merged, _) in enumerate(sheets, start=1):
            rows: dict[int, list[str]] = {}
            for (r, c), v in sorted(cells.items()):
                if isinstance(v, str):
                    strings.append(v)
                    rows.setdefault(r, []).append(f'<c r="{_col(c)}{r}" t="s"><v>{len(strings) - 1}</v></c>')
                else:
                    rows.setdefault(r, []).append(f'<c r="{_col(c)}{r}"><v>{v}</v></c>')
            data = "".join(f'<row r="{r}">{"".join(cs)}</row>' for r, cs in sorted(rows.items()))
            merges = f'<mergeCells count="{len(merged)}">{"".join(f"<mergeCell ref=\"{m}\"/>" for m in merged)}</mergeCells>' if merged else ""
            z.writestr(f"xl/worksheets/sheet{i}.xml", f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{main}">'
                                                       f"<sheetData>{data}</sheetData>{merges}</worksheet>")
        z.writestr("xl/sharedStrings.xml", f'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="{main}">'
                                           + "".join(f"<si><t>{esc(t)}</t></si>" for t in strings) + "</sst>")
        z.writestr("xl/workbook.xml", f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="{main}" xmlns:r="{rel}"><sheets>'
                   + "".join(f'<sheet name="{esc(n)}" sheetId="{i}" r:id="rId{i}"{" state=\"hidden\"" if h else ""}/>'
                             for i, (n, _, _, h) in enumerate(sheets, start=1)) + "</sheets></workbook>")
        z.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships '
                   'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(f'<Relationship Id="rId{i}" Type="{rel}/worksheet" Target="worksheets/sheet{i}.xml"/>'
                             for i in range(1, len(sheets) + 1)) + "</Relationships>")
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types '
                   'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="xml" '
                   'ContentType="application/xml"/></Types>')
    return out.getvalue()


def sealer_xlsx() -> bytes:
    return xlsx([("OCAP - Sealing", SEALING, ["A1:F1", "A4:A5", "B4:B5", "C4:C5", "D4:D5", "F4:F5"], False),
                 ("Troubleshooting", TROUBLESHOOTING, ["A1:E1", "A2:E2", "A9:E9"], False),
                 ("Lists", {(1, 1): "Vertical", (1, 2): "Bottom", (1, 3): "Top"}, [], True)])


SEALING_FIL = {**SEALING, (1, 1): "Plano kapag Wala sa Kontrol", (3, 2): "Pangyayari (Phenomena)",
               (3, 3): "Bahagi ng Makina (Assembly)", (3, 4): "Pangalan sa Centerline (Centerline Name)",
               (3, 5): "Pagsusuri (Inspection)", (3, 6): "Standard na Halaga (Standard Value)",
               (4, 2): "Mahinang seal sa vertical na gilid", (4, 5): "1. Ikumpara ang mga vertical heater sa HMI",
               (5, 5): "2. Tingnan ang lapad ng seal", (6, 2): "Mahinang bottom seal", (6, 5): "1. Tingnan ang bottom jaw",
               (7, 2): "Mahinang top seal", (7, 5): "1. Tingnan ang top jaw"}
TROUBLESHOOTING_FIL = {**TROUBLESHOOTING, (3, 2): "Sealer / Bahagi (Sealer / Area)",
                       (3, 3): "Posibleng Pangyayari (Possible Phenomenon)", (3, 4): "Posibleng Sanhi (Possible Cause)",
                       (3, 5): "Plano ng Operator (Operator Mitigation)",
                       (4, 3): "Madaling matanggal ang seal", (4, 5): "Ihiwalay ang mga pouch at hayaang umabot sa preset ang mga heater.",
                       (5, 3): "Sunog o malutong na seal", (6, 3): "Hindi nakagitna ang seal sa cutter",
                       (7, 3): "Kulang sa 6 mm ang lapad ng vertical seal",
                       (9, 1): "PAALALA: sabihan ang Cell Lead bago patakbuhin nang labas sa centerline."}


def sealer_xlsx_fil() -> bytes:
    """`sealer_xlsx()` in Tagalog, the way the plant would translate it: the same sheets and rows (ADR-0044)."""
    return xlsx([("OCAP - Sealing", SEALING_FIL, ["A1:F1", "A4:A5", "B4:B5", "C4:C5", "D4:D5", "F4:F5"], False),
                 ("Troubleshooting", TROUBLESHOOTING_FIL, ["A1:E1", "A2:E2", "A9:E9"], False),
                 ("Lists", {(1, 1): "Vertical", (1, 2): "Bottom", (1, 3): "Top"}, [], True)])
